"""核验静态学业分析模板下拉约束及分析入口对非法性别的拒绝。"""
import importlib.util
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl

root = Path(__file__).resolve().parents[1]
template = root / "assets" / "academic-analysis-field-template.xlsx"
book = openpyxl.load_workbook(template)
sheet = book.active
assert sheet.max_column == 14 and sheet.max_row == 1
headers = [cell.value for cell in sheet[1]]
assert headers[4] == "性别"
rules = list(sheet.data_validations.dataValidation)
assert len(rules) == 1
rule = rules[0]
assert rule.type == "list" and rule.errorStyle == "stop"
assert rule.showErrorMessage and not rule.showDropDown
assert rule.formula1 == '"男,女"' and str(rule.sqref) == "E2:E1048576"

spec = importlib.util.spec_from_file_location("academic_engine", root / "scripts" / "academic-analysis-engine.py")
engine = importlib.util.module_from_spec(spec)
spec.loader.exec_module(engine)
with TemporaryDirectory(prefix="shuzhi-static-template-qa-") as directory:
    data = openpyxl.Workbook()
    data.active.append(headers)
    for identity, gender in [("QA001", "男"), ("QA002", "女"), ("QA003", "未知")]:
        row = {"学号": identity, "姓名": identity, "性别": gender, "班级": "测试班", "年级": "2026", "课程名称": "测试课程", "总评成绩": "80", "课程性质": "专业必修"}
        data.active.append([row.get(header, "") for header in headers])
    path = Path(directory) / "input.xlsx"
    data.save(path)
    _, _, records, errors = engine.load_records(path)
    assert len(records) == 2 and len(errors) == 1
    assert errors[0]["row"] == 4 and errors[0]["field"] == "性别"
print("Static template QA passed: gender dropdown, Stop alert, valid values accepted and invalid import rejected.")
