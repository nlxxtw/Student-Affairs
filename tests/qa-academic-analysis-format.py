"""Compare a generated academic-analysis workbook with the supplied layout template."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl


SHEET1 = "二级学院考试挂科率统计表"
SHEET2 = "学业成绩班级课程分析统计表"
SHEET3 = "二级学院不及格学生信息"


def same_style(actual, expected, message):
    assert actual._style == expected._style, message


def same_dimensions(actual, expected, columns):
    for column in columns:
        assert actual.column_dimensions[column].width == expected.column_dimensions[column].width, f"{actual.title} {column}列宽不一致"


def same_page(actual, expected):
    for attr in ("orientation", "paperSize", "scale", "fitToWidth", "fitToHeight"):
        assert getattr(actual.page_setup, attr) == getattr(expected.page_setup, attr), f"{actual.title} 页面设置 {attr} 不一致"
    for attr in ("left", "right", "top", "bottom", "header", "footer"):
        assert getattr(actual.page_margins, attr) == getattr(expected.page_margins, attr), f"{actual.title} 页边距 {attr} 不一致"


def verify(reference_path: Path, generated_path: Path):
    reference = openpyxl.load_workbook(reference_path, data_only=False)
    generated = openpyxl.load_workbook(generated_path, data_only=False)
    assert generated.sheetnames == reference.sheetnames, "工作表名称或顺序不一致"

    r1, g1 = reference[SHEET1], generated[SHEET1]
    same_dimensions(g1, r1, "ABCDEFGHIJKLMNOP")
    same_page(g1, r1)
    for row in range(1, 5):
        assert g1.row_dimensions[row].height == r1.row_dimensions[row].height, f"sheet1 第{row}行行高不一致"
        for col in range(1, 16):
            same_style(g1.cell(row, col), r1.cell(row, col), f"sheet1 表头 {row},{col} 样式不一致")
    class_rows = [row for row in range(5, g1.max_row + 1) if isinstance(g1.cell(row, 1).value, int)]
    assert class_rows, "sheet1 没有班级数据"
    for row in class_rows[:-1]:
        for col in range(1, 16):
            same_style(g1.cell(row, col), r1.cell(5, col), f"sheet1 数据行 {row},{col} 样式不一致")
        assert g1.row_dimensions[row].height == 20.15, f"sheet1 第{row}行行高不一致"
    for col in range(1, 16):
        same_style(g1.cell(class_rows[-1], col), r1.cell(11, col), f"sheet1 末数据行 {class_rows[-1]},{col} 样式不一致")
    note_merge = next((item for item in g1.merged_cells.ranges if item.min_col == 1 and item.max_col == 15 and item.min_row >= 5), None)
    assert note_merge is not None and note_merge.max_row == note_merge.min_row + 1, "sheet1 注释合并区域不正确"
    same_style(g1.cell(note_merge.min_row, 1), r1.cell(23, 1), "sheet1 注释样式不一致")
    assert g1.row_dimensions[note_merge.max_row].height == r1.row_dimensions[24].height, "sheet1 注释行高不一致"

    r2, g2 = reference[SHEET2], generated[SHEET2]
    same_dimensions(g2, r2, "ABCDEF")
    same_page(g2, r2)
    title_rows = sorted(item.min_row for item in g2.merged_cells.ranges if item.min_col == 1 and item.max_col == 5 and item.min_row == item.max_row)
    assert len(title_rows) == len(class_rows), "sheet2 班级区块数与sheet1不一致"
    for block_index, title_row in enumerate(title_rows):
        next_title = title_rows[block_index + 1] if block_index + 1 < len(title_rows) else g2.max_row + 1
        same_style(g2.cell(title_row, 1), r2.cell(1, 1), f"sheet2 第{block_index + 1}个标题样式不一致")
        assert g2.row_dimensions[title_row].height == r2.row_dimensions[1].height, f"sheet2 第{block_index + 1}个标题行高不一致"
        for offset, source_row in ((1, 2), (2, 3), (3, 4)):
            assert g2.row_dimensions[title_row + offset].height == r2.row_dimensions[source_row].height, f"sheet2 第{block_index + 1}个区块行高不一致"
            for col in range(1, 6):
                same_style(g2.cell(title_row + offset, col), r2.cell(source_row, col), f"sheet2 第{block_index + 1}个区块表头样式不一致")
        total_row = next((row for row in range(title_row + 4, next_title) if g2.cell(row, 1).value == "合计"), None)
        assert total_row is not None, f"sheet2 第{block_index + 1}个区块缺少合计行"
        for row in range(title_row + 4, total_row):
            assert g2.row_dimensions[row].height == r2.row_dimensions[5].height, f"sheet2 课程行 {row} 行高不一致"
            for col in range(1, 6):
                same_style(g2.cell(row, col), r2.cell(5, col), f"sheet2 课程行 {row},{col} 样式不一致")
        for col in range(1, 6):
            same_style(g2.cell(total_row, col), r2.cell(23, col), f"sheet2 合计行 {total_row},{col} 样式不一致")
        assert g2.row_dimensions[total_row].height == r2.row_dimensions[23].height, f"sheet2 合计行 {total_row} 行高不一致"

    r3, g3 = reference[SHEET3], generated[SHEET3]
    same_dimensions(g3, r3, "ABCDEFGHIJKLM")
    same_page(g3, r3)
    assert {str(item) for item in g3.merged_cells.ranges} == {"A1:M1"}, "sheet3 合并区域不正确"
    for row in (1, 2):
        assert g3.row_dimensions[row].height == r3.row_dimensions[row].height, f"sheet3 第{row}行行高不一致"
        for col in range(1, 14):
            same_style(g3.cell(row, col), r3.cell(row, col), f"sheet3 表头 {row},{col} 样式不一致")
    generated_failures = [row for row in range(3, g3.max_row + 1) if g3.cell(row, 1).value is not None]
    assert generated_failures, "sheet3 没有不及格明细"
    for row in generated_failures:
        for col in range(1, 14):
            same_style(g3.cell(row, col), r3.cell(20, col), f"sheet3 明细 {row},{col} 样式不一致")

    print(f"qa-academic-analysis-format: passed ({len(class_rows)} classes, {len(generated_failures)} failure rows)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: qa-academic-analysis-format.py REFERENCE GENERATED")
    verify(Path(sys.argv[1]), Path(sys.argv[2]))
