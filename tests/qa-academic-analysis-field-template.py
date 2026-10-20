"""Verify that the downloadable academic-analysis template retains only source headers."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl


def verify(source_path: Path, template_path: Path) -> None:
    source = openpyxl.load_workbook(source_path, data_only=False)
    template = openpyxl.load_workbook(template_path, data_only=False)
    assert template.sheetnames == source.sheetnames, "工作表名称或顺序发生变化"
    for source_sheet, template_sheet in zip(source.worksheets, template.worksheets):
        assert template_sheet.max_row == 1, f"{template_sheet.title} 不应包含表头之外的数据行"
        assert template_sheet.max_column == source_sheet.max_column, f"{template_sheet.title} 列数发生变化"
        for column in range(1, source_sheet.max_column + 1):
            source_header = source_sheet.cell(1, column)
            template_header = template_sheet.cell(1, column)
            assert template_header.value == source_header.value, f"{template_sheet.title} 第{column}列表头发生变化"
            assert template_header._style == source_header._style, f"{template_sheet.title} 第{column}列表头格式发生变化"
            assert template_sheet.column_dimensions[template_header.column_letter].width == source_sheet.column_dimensions[source_header.column_letter].width, f"{template_sheet.title} 第{column}列列宽发生变化"
        assert template_sheet.row_dimensions[1].height == source_sheet.row_dimensions[1].height, f"{template_sheet.title} 表头行高发生变化"
    print(f"qa-academic-analysis-field-template: passed ({len(template.sheetnames)} sheet, headers only)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: qa-academic-analysis-field-template.py SOURCE TEMPLATE")
    verify(Path(sys.argv[1]), Path(sys.argv[2]))
