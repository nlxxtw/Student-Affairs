"""回归验证：学业分析成品按身份证地址码填写省市两级生源地。"""
from __future__ import annotations

import importlib.util
import os
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from docx import Document


ROOT = Path(__file__).resolve().parent.parent
ENGINE_PATH = ROOT / "scripts" / "academic-analysis-engine.py"


def load_engine(run_root: Path):
    os.environ["ACADEMIC_ANALYSIS_RUNS"] = str(run_root)
    spec = importlib.util.spec_from_file_location("academic_analysis_engine_qa", ENGINE_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


def build_input(path: Path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    headers = ["学年", "学期", "学号", "姓名", "性别", "课程名称", "课程性质", "总评成绩", "补考成绩", "学院", "年级", "专业", "班级", "身份证号"]
    sheet.append(headers)
    common = ["2025-2026", "2", None, None, "女", "大学英语", "专业必修", 50, "", "教育学院", "2025", "教育学", "852304通信"]
    sheet.append(common[:2] + ["250001", "学生甲"] + common[4:] + ["320102200701011234"])
    sheet.append(common[:2] + ["250002", "学生乙"] + common[4:] + ["110101070101001"])
    sheet.append(common[:2] + ["250003", "学生丙"] + common[4:] + ["990101200701011234"])
    sheet.append(common[:2] + ["250004", "学生丁"] + common[4:] + ["330106200701011234"])
    sheet.append(common[:2] + ["250005", "学生戊"] + common[4:] + ["450102200701011234"])
    sheet.append(common[:2] + ["250006", "学生己"] + common[4:] + ["419001200701011234"])
    workbook.save(path)


def build_teachers(path: Path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["班级", "班主任"])
    sheet.append(["852304", "李老师"])
    workbook.save(path)


def main():
    with tempfile.TemporaryDirectory(prefix="academic-source-qa-") as folder:
        temp = Path(folder)
        source = temp / "input.xlsx"
        teachers = temp / "班主任名单.xlsx"
        build_input(source)
        build_teachers(teachers)
        engine = load_engine(temp / "runs")
        result = engine.process(source.read_bytes(), source.name, teachers, teachers.name)
        output = temp / "runs" / result["runId"] / "学业考试成绩分析统计表.xlsx"
        assert output.is_file(), "统计表未生成"
        workbook = openpyxl.load_workbook(output, data_only=False)
        sheet = workbook["二级学院不及格学生信息"]
        values = {str(sheet.cell(row, 6).value): sheet.cell(row, 9).value for row in range(3, sheet.max_row + 1)}
        assert values["250001"] == "江苏省南京市", values
        assert values["250002"] == "北京市东城区", values
        assert values["250003"] is None, values
        assert values["250004"] == "浙江省杭州市", values
        assert values["250005"] == "广西壮族自治区南宁市", values
        assert values["250006"] == "河南省济源市", values
        assert sheet.cell(2, 9).value == "生源地"
        summary_sheet = workbook["二级学院考试挂科率统计表"]
        assert summary_sheet.cell(5, 6).value == "李老师", "统计表未写入班主任"
        assert sheet.cell(3, 13).value == "李老师", "不及格学生表未写入班主任"
        audit = result["overview"]
        assert audit["teacher_list"]["provided"] is True, audit
        assert audit["teacher_list"]["short_code_matched_classes"] == ["852304通信"], audit
        report_zip = temp / "runs" / result["runId"] / "学业分析报告.zip"
        with zipfile.ZipFile(report_zip) as archive:
            report_name = next(name for name in archive.namelist() if name.endswith(".docx"))
            report_path = temp / "report.docx"
            report_path.write_bytes(archive.read(report_name))
        report = Document(report_path)
        report_text = "\n".join(paragraph.text for paragraph in report.paragraphs)
        assert "班主任：李老师" in report_text, "班级报告未写入班主任"
        assert "【" not in report_text and "】" not in report_text, "班级报告仍含人工补充提示"
        for removed in ("五、重点关注学生", "六、存在问题", "八、总结"):
            assert removed not in report_text, f"班级报告仍含已删除章节：{removed}"
        assert "五、整改措施" in report_text, "精简后的章节序号未连续更新"
    print("qa-academic-analysis-source: passed")


if __name__ == "__main__":
    main()
