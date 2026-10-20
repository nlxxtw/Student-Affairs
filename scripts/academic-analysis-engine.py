"""学业分析助手：本地单文件网页服务。"""
from __future__ import annotations

import hashlib
import json
import math
import os
import shutil
import sys
import tempfile
import threading
import uuid
import zipfile
from copy import copy
from collections import Counter, defaultdict
from datetime import datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote, urlparse
import re

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
STATIC = ROOT / "static"
XLSX_TEMPLATE = ROOT.parent / "assets" / "academic-analysis-template.xlsx"
RUNS = Path(os.environ.get("ACADEMIC_ANALYSIS_RUNS", str(ROOT / "runs"))).resolve()
RUNS.mkdir(exist_ok=True)
REQUIRED = ["学年", "学期", "学号", "姓名", "性别", "课程名称", "课程性质", "总评成绩", "补考成绩", "学院", "年级", "专业", "班级"]
OPTIONAL = ["身份证号"]
EXCLUDED = {"实践奖励", "通识教育选择性必修"}
IDCARD_SOURCE = {
    "11": "北京市", "12": "天津市", "13": "河北省", "14": "山西省", "15": "内蒙古自治区",
    "21": "辽宁省", "22": "吉林省", "23": "黑龙江省",
    "31": "上海市", "32": "江苏省", "33": "浙江省", "34": "安徽省", "35": "福建省", "36": "江西省", "37": "山东省",
    "41": "河南省", "42": "湖北省", "43": "湖南省", "44": "广东省", "45": "广西壮族自治区", "46": "海南省",
    "50": "重庆市", "51": "四川省", "52": "贵州省", "53": "云南省", "54": "西藏自治区",
    "61": "陕西省", "62": "甘肃省", "63": "青海省", "64": "宁夏回族自治区", "65": "新疆维吾尔自治区",
    "71": "台湾省", "81": "香港特别行政区", "82": "澳门特别行政区",
}
IDCARD_CITY = json.loads(r'''{"1301":"石家庄市","1302":"唐山市","1303":"秦皇岛市","1304":"邯郸市","1305":"邢台市","1306":"保定市","1307":"张家口市","1308":"承德市","1309":"沧州市","1310":"廊坊市","1311":"衡水市","1401":"太原市","1402":"大同市","1403":"阳泉市","1404":"长治市","1405":"晋城市","1406":"朔州市","1407":"晋中市","1408":"运城市","1409":"忻州市","1410":"临汾市","1411":"吕梁市","1501":"呼和浩特市","1502":"包头市","1503":"乌海市","1504":"赤峰市","1505":"通辽市","1506":"鄂尔多斯市","1507":"呼伦贝尔市","1508":"巴彦淖尔市","1509":"乌兰察布市","1522":"兴安盟","1525":"锡林郭勒盟","1529":"阿拉善盟","2101":"沈阳市","2102":"大连市","2103":"鞍山市","2104":"抚顺市","2105":"本溪市","2106":"丹东市","2107":"锦州市","2108":"营口市","2109":"阜新市","2110":"辽阳市","2111":"盘锦市","2112":"铁岭市","2113":"朝阳市","2114":"葫芦岛市","2201":"长春市","2202":"吉林市","2203":"四平市","2204":"辽源市","2205":"通化市","2206":"白山市","2207":"松原市","2208":"白城市","2224":"延边朝鲜族自治州","2301":"哈尔滨市","2302":"齐齐哈尔市","2303":"鸡西市","2304":"鹤岗市","2305":"双鸭山市","2306":"大庆市","2307":"伊春市","2308":"佳木斯市","2309":"七台河市","2310":"牡丹江市","2311":"黑河市","2312":"绥化市","2327":"大兴安岭地区","3201":"南京市","3202":"无锡市","3203":"徐州市","3204":"常州市","3205":"苏州市","3206":"南通市","3207":"连云港市","3208":"淮安市","3209":"盐城市","3210":"扬州市","3211":"镇江市","3212":"泰州市","3213":"宿迁市","3301":"杭州市","3302":"宁波市","3303":"温州市","3304":"嘉兴市","3305":"湖州市","3306":"绍兴市","3307":"金华市","3308":"衢州市","3309":"舟山市","3310":"台州市","3311":"丽水市","3401":"合肥市","3402":"芜湖市","3403":"蚌埠市","3404":"淮南市","3405":"马鞍山市","3406":"淮北市","3407":"铜陵市","3408":"安庆市","3410":"黄山市","3411":"滁州市","3412":"阜阳市","3413":"宿州市","3415":"六安市","3416":"亳州市","3417":"池州市","3418":"宣城市","3501":"福州市","3502":"厦门市","3503":"莆田市","3504":"三明市","3505":"泉州市","3506":"漳州市","3507":"南平市","3508":"龙岩市","3509":"宁德市","3601":"南昌市","3602":"景德镇市","3603":"萍乡市","3604":"九江市","3605":"新余市","3606":"鹰潭市","3607":"赣州市","3608":"吉安市","3609":"宜春市","3610":"抚州市","3611":"上饶市","3701":"济南市","3702":"青岛市","3703":"淄博市","3704":"枣庄市","3705":"东营市","3706":"烟台市","3707":"潍坊市","3708":"济宁市","3709":"泰安市","3710":"威海市","3711":"日照市","3713":"临沂市","3714":"德州市","3715":"聊城市","3716":"滨州市","3717":"菏泽市","4101":"郑州市","4102":"开封市","4103":"洛阳市","4104":"平顶山市","4105":"安阳市","4106":"鹤壁市","4107":"新乡市","4108":"焦作市","4109":"濮阳市","4110":"许昌市","4111":"漯河市","4112":"三门峡市","4113":"南阳市","4114":"商丘市","4115":"信阳市","4116":"周口市","4117":"驻马店市","4201":"武汉市","4202":"黄石市","4203":"十堰市","4205":"宜昌市","4206":"襄阳市","4207":"鄂州市","4208":"荆门市","4209":"孝感市","4210":"荆州市","4211":"黄冈市","4212":"咸宁市","4213":"随州市","4228":"恩施土家族苗族自治州","4301":"长沙市","4302":"株洲市","4303":"湘潭市","4304":"衡阳市","4305":"邵阳市","4306":"岳阳市","4307":"常德市","4308":"张家界市","4309":"益阳市","4310":"郴州市","4311":"永州市","4312":"怀化市","4313":"娄底市","4331":"湘西土家族苗族自治州","4401":"广州市","4402":"韶关市","4403":"深圳市","4404":"珠海市","4405":"汕头市","4406":"佛山市","4407":"江门市","4408":"湛江市","4409":"茂名市","4412":"肇庆市","4413":"惠州市","4414":"梅州市","4415":"汕尾市","4416":"河源市","4417":"阳江市","4418":"清远市","4419":"东莞市","4420":"中山市","4451":"潮州市","4452":"揭阳市","4453":"云浮市","4501":"南宁市","4502":"柳州市","4503":"桂林市","4504":"梧州市","4505":"北海市","4506":"防城港市","4507":"钦州市","4508":"贵港市","4509":"玉林市","4510":"百色市","4511":"贺州市","4512":"河池市","4513":"来宾市","4514":"崇左市","4601":"海口市","4602":"三亚市","4603":"三沙市","4604":"儋州市","5101":"成都市","5103":"自贡市","5104":"攀枝花市","5105":"泸州市","5106":"德阳市","5107":"绵阳市","5108":"广元市","5109":"遂宁市","5110":"内江市","5111":"乐山市","5113":"南充市","5114":"眉山市","5115":"宜宾市","5116":"广安市","5117":"达州市","5118":"雅安市","5119":"巴中市","5120":"资阳市","5132":"阿坝藏族羌族自治州","5133":"甘孜藏族自治州","5134":"凉山彝族自治州","5201":"贵阳市","5202":"六盘水市","5203":"遵义市","5204":"安顺市","5205":"毕节市","5206":"铜仁市","5223":"黔西南布依族苗族自治州","5226":"黔东南苗族侗族自治州","5227":"黔南布依族苗族自治州","5301":"昆明市","5303":"曲靖市","5304":"玉溪市","5305":"保山市","5306":"昭通市","5307":"丽江市","5308":"普洱市","5309":"临沧市","5323":"楚雄彝族自治州","5325":"红河哈尼族彝族自治州","5326":"文山壮族苗族自治州","5328":"西双版纳傣族自治州","5329":"大理白族自治州","5331":"德宏傣族景颇族自治州","5333":"怒江傈僳族自治州","5334":"迪庆藏族自治州","5401":"拉萨市","5402":"日喀则市","5403":"昌都市","5404":"林芝市","5405":"山南市","5406":"那曲市","5425":"阿里地区","6101":"西安市","6102":"铜川市","6103":"宝鸡市","6104":"咸阳市","6105":"渭南市","6106":"延安市","6107":"汉中市","6108":"榆林市","6109":"安康市","6110":"商洛市","6201":"兰州市","6202":"嘉峪关市","6203":"金昌市","6204":"白银市","6205":"天水市","6206":"武威市","6207":"张掖市","6208":"平凉市","6209":"酒泉市","6210":"庆阳市","6211":"定西市","6212":"陇南市","6229":"临夏回族自治州","6230":"甘南藏族自治州","6301":"西宁市","6302":"海东市","6322":"海北藏族自治州","6323":"黄南藏族自治州","6325":"海南藏族自治州","6326":"果洛藏族自治州","6327":"玉树藏族自治州","6328":"海西蒙古族藏族自治州","6401":"银川市","6402":"石嘴山市","6403":"吴忠市","6404":"固原市","6405":"中卫市","6501":"乌鲁木齐市","6502":"克拉玛依市","6504":"吐鲁番市","6505":"哈密市","6523":"昌吉回族自治州","6527":"博尔塔拉蒙古自治州","6528":"巴音郭楞蒙古自治州","6529":"阿克苏地区","6530":"克孜勒苏柯尔克孜自治州","6531":"喀什地区","6532":"和田地区","6540":"伊犁哈萨克自治州","6542":"塔城地区","6543":"阿勒泰地区","7101":"台北市","7102":"高雄市","7103":"基隆市","7104":"台中市","7105":"台南市","7106":"新竹市","7107":"嘉义市","8100":"離島區","8200":"聖方濟各堂區"}''')
IDCARD_DIRECT_COUNTY = json.loads(r'''{"110101":"北京市东城区","110102":"北京市西城区","110105":"北京市朝阳区","110106":"北京市丰台区","110107":"北京市石景山区","110108":"北京市海淀区","110109":"北京市门头沟区","110111":"北京市房山区","110112":"北京市通州区","110113":"北京市顺义区","110114":"北京市昌平区","110115":"北京市大兴区","110116":"北京市怀柔区","110117":"北京市平谷区","110118":"北京市密云区","110119":"北京市延庆区","120101":"天津市和平区","120102":"天津市河东区","120103":"天津市河西区","120104":"天津市南开区","120105":"天津市河北区","120106":"天津市红桥区","120110":"天津市东丽区","120111":"天津市西青区","120112":"天津市津南区","120113":"天津市北辰区","120114":"天津市武清区","120115":"天津市宝坻区","120116":"天津市滨海新区","120117":"天津市宁河区","120118":"天津市静海区","120119":"天津市蓟州区","310101":"上海市黄浦区","310104":"上海市徐汇区","310105":"上海市长宁区","310106":"上海市静安区","310107":"上海市普陀区","310109":"上海市虹口区","310110":"上海市杨浦区","310112":"上海市闵行区","310113":"上海市宝山区","310114":"上海市嘉定区","310115":"上海市浦东新区","310116":"上海市金山区","310117":"上海市松江区","310118":"上海市青浦区","310120":"上海市奉贤区","310151":"上海市崇明区","419001":"河南省济源市","429004":"湖北省仙桃市","429005":"湖北省潜江市","429006":"湖北省天门市","429021":"湖北省神农架林区","469001":"海南省五指山市","469002":"海南省琼海市","469005":"海南省文昌市","469006":"海南省万宁市","469007":"海南省东方市","469021":"海南省定安县","469022":"海南省屯昌县","469023":"海南省澄迈县","469024":"海南省临高县","469025":"海南省白沙黎族自治县","469026":"海南省昌江黎族自治县","469027":"海南省乐东黎族自治县","469028":"海南省陵水黎族自治县","469029":"海南省保亭黎族苗族自治县","469030":"海南省琼中黎族苗族自治县","500101":"重庆市万州区","500102":"重庆市涪陵区","500103":"重庆市渝中区","500104":"重庆市大渡口区","500105":"重庆市江北区","500106":"重庆市沙坪坝区","500107":"重庆市九龙坡区","500108":"重庆市南岸区","500109":"重庆市北碚区","500110":"重庆市綦江区","500111":"重庆市大足区","500112":"重庆市渝北区","500113":"重庆市巴南区","500114":"重庆市黔江区","500115":"重庆市长寿区","500116":"重庆市江津区","500117":"重庆市合川区","500118":"重庆市永川区","500119":"重庆市南川区","500120":"重庆市璧山区","500151":"重庆市铜梁区","500152":"重庆市潼南区","500153":"重庆市荣昌区","500154":"重庆市开州区","500155":"重庆市梁平区","500156":"重庆市武隆区","500229":"重庆市城口县","500230":"重庆市丰都县","500231":"重庆市垫江县","500233":"重庆市忠县","500235":"重庆市云阳县","500236":"重庆市奉节县","500237":"重庆市巫山县","500238":"重庆市巫溪县","500240":"重庆市石柱土家族自治县","500241":"重庆市秀山土家族苗族自治县","500242":"重庆市酉阳土家族苗族自治县","500243":"重庆市彭水苗族土家族自治县","659001":"新疆维吾尔自治区石河子市","659002":"新疆维吾尔自治区阿拉尔市","659003":"新疆维吾尔自治区图木舒克市","659004":"新疆维吾尔自治区五家渠市","659006":"新疆维吾尔自治区铁门关市"}''')
SCORE_MAP = {"优秀": 90, "良好": 80, "通过": 80, "中等": 70, "及格": 60,
             "不及格": 50, "不合格": 50, "旷考": 0, "作弊": 0,
             "合格": 80, "缺考": 0, "舞弊": 0}
PASS_WORDS = {"优秀", "良好", "通过", "中等", "及格", "合格"}


def clean(v): return "" if v is None else str(v).strip()

def class_key(v): return re.sub(r"[\s\u3000]+", "", clean(v))

def class_code(v):
    """班级前六码仅在其为六位数字时作为安全的短码匹配依据。"""
    key = class_key(v)
    return key[:6] if len(key) >= 6 and key[:6].isdigit() else ""

def source_from_idcard(value):
    """按标准行政区输出两级生源地；直辖市以区县作为第二级。"""
    identity = re.sub(r"[\s-]", "", clean(value)).upper()
    if not (re.fullmatch(r"\d{15}", identity) or re.fullmatch(r"\d{17}[\dX]", identity)):
        return ""
    province_code = identity[:2]
    province = IDCARD_SOURCE.get(province_code, "")
    if not province:
        return ""
    direct = IDCARD_DIRECT_COUNTY.get(identity[:6])
    if direct:
        return direct
    city = IDCARD_CITY.get(identity[:4])
    return province + city if city else ""

def course_key(s):
    return clean(s).replace("（", "(").replace("）", ")").replace(" ", "").replace("\u3000", "")

def number(v):
    s = clean(v)
    try: return float(s)
    except ValueError: return SCORE_MAP.get(s)

def is_fail(v):
    n = number(v)
    return None if n is None else n < 60

def makeup_status(record):
    makeup = number(record.get("补考成绩"))
    if makeup is not None and makeup >= 60:
        return "通过"
    if not clean(record.get("补考成绩")) and clean(record.get("课程性质")) == "自主发展":
        return "未通过无需重修"
    return "未通过待重修"

def pct(v): return f"{v * 100:.2f}%"

def semester_name(v): return {"1": "第一学期", "2": "第二学期", "3": "第三学期"}.get(clean(v), f"第{clean(v)}学期")

def font(size=18, bold=False):
    for p in ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf", "C:/Windows/Fonts/simsun.ttc"):
        if Path(p).exists(): return ImageFont.truetype(p, size, index=0)
    return ImageFont.load_default()

def canvas(title, labels, values, kind, path):
    w, h = 1100, 620; img = Image.new("RGB", (w, h), "white"); d = ImageDraw.Draw(img)
    d.text((42, 26), title, fill="#183153", font=font(28))
    if kind == "pie":
        total = sum(values) or 1; box = (70, 100, 560, 590); start = 0; colors = ["#2F6B9A", "#4D9F93", "#F0A65A", "#D66B5D", "#7D8CC4"]
        for i, (lab, val) in enumerate(zip(labels, values)):
            end = start + val / total * 360; d.pieslice(box, start, end, fill=colors[i % len(colors)], outline="white", width=2); start = end
            d.rectangle((630, 120 + i * 70, 654, 144 + i * 70), fill=colors[i % len(colors)])
            d.text((668, 118 + i * 70), f"{lab}  {val}人 ({val/total*100:.1f}%)", fill="#263746", font=font(18))
    else:
        maxv = max(values) if values else 1; left, base, chartw, charth = 70, 540, 960, 380; gap = 20; bw = max(18, (chartw - gap * (len(values)+1)) / max(len(values), 1))
        d.line((left, base, left + chartw, base), fill="#778899", width=2)
        for i, (lab, val) in enumerate(zip(labels, values)):
            x = left + gap + i * (bw + gap); bh = (val / maxv * charth) if maxv else 0
            d.rectangle((x, base-bh, x+bw, base), fill="#2F6B9A" if kind != "bucket" else "#4D9F93")
            d.text((x, base-bh-28), str(val), fill="#183153", font=font(16))
            d.multiline_text((x, base+10), lab, fill="#263746", font=font(14), spacing=2)
    img.save(path)


def load_records(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active; iterator = ws.iter_rows(values_only=True); headers = [clean(x) for x in next(iterator, ())]
    missing = [x for x in REQUIRED if x not in headers]
    if missing:
        wb.close()
        raise ValueError("缺少必需列：" + "、".join(missing))
    rows = [{h: clean(v) for h, v in zip(headers, raw)} for raw in iterator]
    wb.close()
    unknown = sorted({r["总评成绩"] for r in rows if r["总评成绩"] and number(r["总评成绩"]) is None})
    if unknown: raise ValueError("总评成绩存在待确认文本：" + "、".join(unknown))
    errors = []
    usable = []
    for index, r in enumerate(rows, 2):
        if any(not r[x] for x in ("学号", "班级", "年级", "课程名称", "总评成绩")):
            errors.append({"row": index, "type": "关键字段为空"}); continue
        if r["性别"] and r["性别"] not in {"男", "女"}:
            errors.append({"row": index, "field": "性别", "type": "固定选项错误", "value": r["性别"], "message": "性别只能选择男或女"}); continue
        if r["课程性质"] in EXCLUDED: continue
        r["课程键"] = course_key(r["课程名称"]); usable.append(r)
    seen = set(); records = []
    for r in usable:
        key = (r["学号"], r["课程键"])
        if key in seen:
            errors.append({"student": r["学号"], "course": r["课程名称"], "type": "重复成绩记录，保留首条"}); continue
        seen.add(key); records.append(r)
    return headers, rows, records, errors


def load_teachers(path: Path, class_names):
    """读取可选班主任名单；只接受清晰的班级—班主任对应关系。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator, ())]
    class_column = next((index for index, value in enumerate(headers) if value in {"班级", "班级名称", "行政班"}), None)
    teacher_column = next((index for index, value in enumerate(headers) if value in {"班主任", "班主任姓名", "班主任老师"}), None)
    if class_column is None or teacher_column is None:
        raise ValueError("班主任名单缺少必需列：班级、班主任（可识别别名：班级名称、行政班、班主任姓名）")
    teachers, teacher_sources, issues, duplicate_rows = {}, {}, [], 0
    for row_number, raw in enumerate(iterator, 2):
        values = [clean(value) for value in raw]
        class_name = values[class_column] if class_column < len(values) else ""
        teacher = values[teacher_column] if teacher_column < len(values) else ""
        if not class_name and not teacher:
            continue
        key = class_key(class_name)
        if not key or not teacher:
            issues.append({"source": "班主任名单", "row": row_number, "type": "班级或班主任为空"})
            continue
        if key in teachers:
            duplicate_rows += 1
            issues.append({"source": "班主任名单", "row": row_number, "type": "班级重复，保留首条", "class": class_name})
            continue
        teachers[key] = teacher
        teacher_sources[key] = class_name
    score_classes = {class_key(name): name for name in class_names}
    resolved, used_teacher_keys, short_code_matches = {}, set(), []
    for score_key, score_name in score_classes.items():
        if score_key in teachers:
            resolved[score_key] = teachers[score_key]
            used_teacher_keys.add(score_key)
            continue
        code = class_code(score_key)
        short_keys = [key for key in teachers if key == code]
        same_code_classes = [name for key, name in score_classes.items() if class_code(key) == code]
        if len(short_keys) == 1 and len(same_code_classes) == 1:
            resolved[score_key] = teachers[short_keys[0]]
            used_teacher_keys.add(short_keys[0])
            short_code_matches.append(score_name)
        elif short_keys and len(same_code_classes) > 1:
            issues.append({"source": "班主任名单", "type": "班级前6位匹配存在歧义，未自动匹配", "code": code, "classes": sorted(same_code_classes)})
    extra = sorted({name for key, name in score_classes.items() if key not in resolved})
    outside = sorted({teacher_sources[key] for key in teachers if key not in used_teacher_keys})
    if extra:
        issues.append({"source": "班主任名单", "type": "本次成绩表中的班级未匹配班主任", "classes": extra})
    if outside:
        issues.append({"source": "班主任名单", "type": "名单中的班级未出现在本次成绩表", "classes": [key for key in outside]})
    return resolved, {"provided": True, "rows": len(teachers), "duplicate_rows": duplicate_rows, "matched_classes": len(score_classes) - len(extra), "short_code_matched_classes": short_code_matches, "unmatched_classes": extra, "outside_classes": outside}, issues


def stats_for(records, all_rows):
    by_class = defaultdict(list)
    for r in records: by_class[r["班级"]].append(r)
    raw_students = defaultdict(set)
    for r in all_rows:
        if r.get("班级"): raw_students[r["班级"]].add(r.get("学号"))
    output = []
    for cls, recs in by_class.items():
        meta = recs[0]; courses = defaultdict(list); failures = []
        for r in recs:
            courses[r["课程键"]].append(r)
            if is_fail(r["总评成绩"]): failures.append(r)
        by_student = defaultdict(list)
        for r in failures: by_student[r["学号"]].append(r)
        students = raw_students[cls]
        profile = {}
        for r in all_rows:
            if r.get("班级") == cls: profile.setdefault(r["学号"], r)
        genders = Counter(profile[s].get("性别") for s in students if s in profile)
        bucket = [sum(len(x)==1 for x in by_student.values()), sum(2<=len(x)<=4 for x in by_student.values()), sum(5<=len(x)<=9 for x in by_student.values()), sum(len(x)>=10 for x in by_student.values())]
        crows=[]
        for key, items in sorted(courses.items(), key=lambda z: z[0]):
            fail_n=sum(is_fail(x["总评成绩"]) for x in items)
            crows.append({"name": items[0]["课程名称"], "key": key, "enrolled":len(items), "failed":fail_n, "items":items})
        output.append({"class": cls, "year": meta["年级"], "major": meta["专业"], "college": meta["学院"], "students":len(students), "male":genders["男"], "female":genders["女"], "records":len(recs), "courses":crows, "failures":failures, "fail_students":by_student, "bucket":bucket})
    return sorted(output, key=lambda x: (x["year"], x["class"][:6], x["class"]))


def row_style(ws, row, last_col):
    return [copy(ws.cell(row, col)._style) for col in range(1, last_col + 1)]


def apply_row_style(ws, row, styles, height=None):
    for col, style in enumerate(styles, 1):
        ws.cell(row, col)._style = copy(style)
    ws.row_dimensions[row].height = height


def clear_sheet_values(ws):
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    for row in ws.iter_rows():
        for target in row:
            target.value = None


def set_row_values(ws, row, values):
    for col, value in enumerate(values, 1):
        ws.cell(row, col, value)


def create_xlsx(classes, sample, out):
    if not XLSX_TEMPLATE.is_file():
        raise FileNotFoundError(f"学业分析统计表版式模板缺失：{XLSX_TEMPLATE}")
    wb = openpyxl.load_workbook(XLSX_TEMPLATE)
    s1 = wb["二级学院考试挂科率统计表"]
    s2 = wb["学业成绩班级课程分析统计表"]
    s3 = wb["二级学院不及格学生信息"]
    year, sem = sample["year"], sample["semester"]

    # 第一张表：保留参考表的列宽、页面设置和四行复合表头，数据行按班级数扩展。
    s1_header_styles = {row: row_style(s1, row, 15) for row in range(1, 5)}
    s1_data_style = row_style(s1, 5, 15)
    s1_last_data_style = row_style(s1, 11, 15)
    s1_blank_style = row_style(s1, 12, 15)
    s1_note_styles = {23: row_style(s1, 23, 15), 24: row_style(s1, 24, 15)}
    s1_heights = {row: s1.row_dimensions[row].height for row in range(1, 25)}
    clear_sheet_values(s1)
    for row in range(1, 5):
        apply_row_style(s1, row, s1_header_styles[row], s1_heights[row])
    s1.merge_cells("A1:O1")
    s1.merge_cells("A2:O2")
    for col in range(1, 11):
        s1.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
    s1.merge_cells("K3:O3")
    s1["A1"] = f"南京师范大学中北学院{year}学年{sem}考试挂科率统计表"
    s1["A2"] = "二级学院（公章）：                                学工办主任（签名）：                            分管学工院领导（签名）："
    for col, value in enumerate(["序号", "年级", "班级", "专业", "人数", "班主任", "总考试门次", "总考试人次", "不及格人次", "挂科率"], 1):
        s1.cell(3, col, value)
    s1["K3"] = "挂科人数"
    for col, value in enumerate(["挂科总人数", "1门不及格人数", "2-4门不及格人数", "5-9门不及格人数", "10门及以上不及格人数"], 11):
        s1.cell(4, col, value)
    data_end = 4 + len(classes)
    note_row = max(23, data_end + 1)
    for row in range(5, note_row):
        style = s1_last_data_style if row == data_end and classes else (s1_data_style if row <= data_end else s1_blank_style)
        apply_row_style(s1, row, style, 20.15)
    for number_index, current in enumerate(classes, 1):
        row = 4 + number_index
        set_row_values(s1, row, [number_index, current["year"], current["class"][:6], current["major"], current["students"], current.get("teacher", ""), len(current["courses"]), current["records"], len(current["failures"]), f"=I{row}/H{row}", len(current["fail_students"]), *current["bucket"]])
        s1.cell(row, 10).number_format = "0.00%"
    apply_row_style(s1, note_row, s1_note_styles[23], s1_heights[23])
    apply_row_style(s1, note_row + 1, s1_note_styles[24], s1_heights[24])
    s1.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 1, end_column=15)
    s1.cell(note_row, 1, "注释：\n1.总考试门次=sheet2的每个班的课程数之和；\n2.总考试人次=sheet2的修读人数之和；\n3.不及格人次=sheet2的挂科人数之和；\n4.挂科率=不及格人次/总考试人次。")

    # 第二张表：每个班级都是参考表中的完整块，课程行根据实际课程数动态复制。
    s2_styles = {
        "title": row_style(s2, 1, 5),
        "info1": row_style(s2, 2, 5),
        "info2": row_style(s2, 3, 5),
        "header": row_style(s2, 4, 5),
        "course": row_style(s2, 5, 5),
        "total": row_style(s2, 23, 5),
        "blank1": row_style(s2, 24, 5),
        "blank2": row_style(s2, 25, 5),
    }
    s2_heights = {row: s2.row_dimensions[row].height for row in range(1, 26)}
    clear_sheet_values(s2)
    row = 1
    for current in classes:
        apply_row_style(s2, row, s2_styles["title"], s2_heights[1])
        s2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        s2.cell(row, 1, f"南京师范大学中北学院 {year}学年{sem}\n学业成绩课程分析报表")
        apply_row_style(s2, row + 1, s2_styles["info1"], s2_heights[2])
        set_row_values(s2, row + 1, ["二级学院：", current["college"], f"年级：{current['year']}级", f"专业：{current['major']}", f"班级：{current['class']}"])
        apply_row_style(s2, row + 2, s2_styles["info2"], s2_heights[3])
        set_row_values(s2, row + 2, ["班级人数：", current["students"], f"男生人数：{current['male']}", f"女生人数：{current['female']}", f"班主任：{current.get('teacher', '')}"])
        apply_row_style(s2, row + 3, s2_styles["header"], s2_heights[4])
        set_row_values(s2, row + 3, ["序号", "课程名称", "修读人数", "挂科人数", "挂科率"])
        course_start = row + 4
        for course_index, course in enumerate(current["courses"], 1):
            course_row = course_start + course_index - 1
            apply_row_style(s2, course_row, s2_styles["course"], s2_heights[5])
            set_row_values(s2, course_row, [course_index, course["name"], course["enrolled"], course["failed"], f"=D{course_row}/C{course_row}"])
            s2.cell(course_row, 5).number_format = "0.00%"
        total_row = course_start + len(current["courses"])
        apply_row_style(s2, total_row, s2_styles["total"], s2_heights[23])
        set_row_values(s2, total_row, ["合计", len(current["courses"]), f"=SUM(C{course_start}:C{total_row - 1})", f"=SUM(D{course_start}:D{total_row - 1})", f"=D{total_row}/C{total_row}"])
        s2.cell(total_row, 5).number_format = "0.00%"
        apply_row_style(s2, total_row + 1, s2_styles["blank1"], s2_heights[24])
        apply_row_style(s2, total_row + 2, s2_styles["blank2"], s2_heights[25])
        row = total_row + 3

    # 第三张表：标题、表头、列宽和打印设置取自参考表，明细行只写本次分析数据。
    s3_title_style = row_style(s3, 1, 13)
    s3_header_style = row_style(s3, 2, 13)
    s3_data_style = row_style(s3, 20, 13)
    s3_title_height = s3.row_dimensions[1].height
    s3_header_height = s3.row_dimensions[2].height
    s3_data_height = s3.row_dimensions[20].height
    clear_sheet_values(s3)
    apply_row_style(s3, 1, s3_title_style, s3_title_height)
    apply_row_style(s3, 2, s3_header_style, s3_header_height)
    s3.merge_cells("A1:M1")
    s3["A1"] = f"南京师范大学中北学院{year}学年{sem}班主任所带班级不及格学生信息统计表"
    set_row_values(s3, 2, ["序号", "二级学院", "年级", "班级", "专业", "学号", "学生姓名", "性别", "生源地", "不及格科目", "不及格科目原始成绩", "补考情况", "班主任"])
    failure_index = 0
    for current in classes:
        for failure in sorted(current["failures"], key=lambda item: (item["学号"], item["课程键"])):
            failure_index += 1
            failure_row = failure_index + 2
            apply_row_style(s3, failure_row, s3_data_style, s3_data_height)
            set_row_values(s3, failure_row, [failure_index, failure["学院"], failure["年级"], failure["班级"], failure["专业"], failure["学号"], failure["姓名"], failure["性别"], source_from_idcard(failure.get("身份证号")), failure["课程名称"], failure["总评成绩"], makeup_status(failure), current.get("teacher", "")])

    wb.save(out)


def doc_font(run, name="宋体", size=11, bold=False):
    run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run._element.rPr.rFonts.set(qn("w:eastAsia"),name)
def add_table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers));t.style="Table Grid"
    for i,h in enumerate(headers): doc_font(t.cell(0,i).paragraphs[0].add_run(str(h)),"黑体",9,True)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):doc_font(cells[i].paragraphs[0].add_run(str(v)),"宋体",8.5)
    return t
def heading(doc,text):
    p=doc.add_paragraph();doc_font(p.add_run(text),"黑体",14,True);return p
def para(doc,text):
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(24);p.paragraph_format.line_spacing=1.3;p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;doc_font(p.add_run(text));return p

def caption(doc, text):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;doc_font(p.add_run(text),"宋体",9);return p

def create_report(x, sample, folder):
    doc=Document(); p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;doc_font(p.add_run(f"{x['class']}{sample['year']}学年{sample['semester']}学业分析报告"),"黑体",18,True)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;doc_font(p.add_run("班主任："+(x.get("teacher") or "未提供")))
    heading(doc,"一、班级基本情况");add_table(doc,["项目","内容"],[["二级学院",x["college"]],["年级",x["year"]+"级"],["专业",x["major"]],["班级",x["class"]],["班级人数",f"{x['students']}人（男{x['male']}人、女{x['female']}人）"],["总考试门次",len(x['courses'])],["总考试人次",x['records']]])
    para(doc,"班级整体目标：在准确掌握课程通过情况和学生成绩结构的基础上，实施分层学业指导，持续降低不及格风险并提升整体学习质量。")
    rate=len(x['failures'])/x['records'] if x['records'] else 0; fail_counts=Counter(len(v) for v in x['fail_students'].values()); max_fail=max(fail_counts) if fail_counts else 0
    makeup_pass=sum(makeup_status(r)=="通过" for r in x['failures']); pending=sum(makeup_status(r)=="未通过待重修" for r in x['failures']); no_retake=sum(makeup_status(r)=="未通过无需重修" for r in x['failures'])
    pending_courses=Counter(r['课程名称'] for r in x['failures'] if makeup_status(r)=="未通过待重修")
    max_fail_students=sum(len(v)==max_fail for v in x['fail_students'].values()) if max_fail else 0
    heading(doc,"二、考试情况概述")
    if x['failures']:
        para(doc,f"本班本学期共统计{len(x['courses'])}门课程、{x['records']}人次成绩，不及格{len(x['failures'])}人次，班级挂科率为{pct(rate)}。共有{len(x['fail_students'])}名学生出现不及格，占本班人数的{pct(len(x['fail_students'])/(x['students'] or 1))}；单人最多不及格{max_fail}门，涉及{max_fail_students}人。")
        makeup_text=f"补考已通过{makeup_pass}人次，补考通过率为{pct(makeup_pass/len(x['failures']))}；补考后仍待重修{pending}人次"
        if no_retake: makeup_text+=f"，另有{no_retake}人次按课程性质判定为无需重修"
        if pending_courses:
            course_text="、".join(f"{name}{count}人次" for name,count in pending_courses.most_common())
            makeup_text+=f"。补考后仍未通过的课程包括：{course_text}"
        para(doc,makeup_text+"。")
    else:
        para(doc,f"本班本学期共统计{len(x['courses'])}门课程、{x['records']}人次成绩，全部课程通过，班级挂科率为0.00%。本学期无需安排补考或重修跟踪。")
    scores=defaultdict(list)
    for c in x['courses']:
        for r in c['items']:
            if number(r['总评成绩']) is not None:scores[r['学号']].append(number(r['总评成绩']))
    avgs=[sum(v)/len(v) for v in scores.values() if v]; avg=sum(avgs)/len(avgs) if avgs else 0; sorted_a=sorted(avgs); med=sorted_a[len(sorted_a)//2] if len(sorted_a)%2 else (sorted_a[len(sorted_a)//2-1]+sorted_a[len(sorted_a)//2])/2 if sorted_a else 0
    lowest=min(avgs) if avgs else 0; highest=max(avgs) if avgs else 0; deviation=math.sqrt(sum((v-avg)**2 for v in avgs)/len(avgs)) if avgs else 0; pass_rate=sum(v>=60 for v in avgs)/(len(avgs) or 1)
    seg=[sum(a>=90 for a in avgs),sum(80<=a<90 for a in avgs),sum(70<=a<80 for a in avgs),sum(60<=a<70 for a in avgs),sum(a<60 for a in avgs)]
    chartdir=folder/"charts";chartdir.mkdir(exist_ok=True); charts=[chartdir/"chart1.png",chartdir/"chart2.png",chartdir/"chart3.png",chartdir/"chart4.png"]
    top=sorted([c for c in x['courses'] if c['failed']],key=lambda c:(-c['failed']/c['enrolled'],-c['failed'],c['key']))[:10]
    canvas("报告图1 学生平均分分布",["90+","80-89","70-79","60-69","<60"],seg,"bar",charts[0]);canvas("报告图2 课程挂科率（%）",[f"{c['name'][:7]}\n({c['failed']}人)" for c in top],[round(c['failed']/c['enrolled']*100,2) for c in top],"bar",charts[1]);canvas("报告图3 不及格门数分布",["1门","2-4门","5-9门","10+"],x['bucket'],"bucket",charts[2]);canvas("报告图4 成绩分数段占比",["90+","80-89","70-79","60-69","<60"],seg,"pie",charts[3])
    heading(doc,"三、成绩水平分析");para(doc,f"按每名学生过滤后全部课程总评成绩的算术平均分统计，班级平均分为{avg:.2f}分，中位数为{med:.2f}分，最高分为{highest:.2f}分，最低分为{lowest:.2f}分，标准差为{deviation:.2f}，学生平均分及格率为{pct(pass_rate)}。分数段分布如下。")
    add_table(doc,["分数段","人数","占比"],[[k,v,pct(v/(len(avgs) or 1))] for k,v in zip(["90分及以上（优秀）","80-89分（良好）","70-79分（中等）","60-69分（及格）","60分以下（不及格）"],seg)])
    good=seg[0]+seg[1]; middle=seg[2]; lower=seg[3]+seg[4]
    spread="较为集中" if deviation<5 else ("存在一定梯度" if deviation<=10 else "差异较为明显")
    relation="高于" if med>avg+0.01 else ("低于" if med<avg-0.01 else "接近")
    feature=f"成绩分布特征：80分及以上学生{good}人，占{pct(good/(len(avgs) or 1))}，构成班级学业骨干；70—79分学生{middle}人，占{pct(middle/(len(avgs) or 1))}；70分以下学生{lower}人，占{pct(lower/(len(avgs) or 1))}。中位数{relation}平均分，成绩标准差为{deviation:.2f}，说明班级学生成绩{spread}。"
    if seg[0]: feature+=f"班级已有{seg[0]}名学生进入90分以上高分段，可发挥示范带动作用。"
    if seg[3] or seg[4]: feature+=f"其中60—69分临界段{seg[3]}人、平均分低于60分{seg[4]}人，应纳入分层预警和持续跟踪。"
    para(doc,feature)
    doc.add_picture(str(charts[0]),width=Cm(15));caption(doc,"报告图1 学生平均分分布")
    doc.add_picture(str(charts[3]),width=Cm(15));caption(doc,"报告图4 成绩分数段占比")
    heading(doc,"四、课程分析")
    course_rows=[]
    course_analysis=sorted(x['courses'],key=lambda c:(-c['failed'],-c['failed']/c['enrolled'],c['key']))
    for c in course_analysis:
        passed=sum(makeup_status(r)=="通过" for r in c['items'] if is_fail(r['总评成绩']))
        still=sum(makeup_status(r)=="未通过待重修" for r in c['items'] if is_fail(r['总评成绩']))
        course_rows.append([c['name'],c['enrolled'],c['failed'],pct(c['failed']/c['enrolled']),passed,still,"整班课程" if c['enrolled']>=x['students']*0.5 else "重修/个别修读"])
    add_table(doc,["课程名称","修读人数","挂科人数","挂科率","补考通过","待重修","课程类型"],course_rows)
    whole_risk=[c for c in course_analysis if c['failed'] and c['enrolled']>=x['students']*0.5]; individual=[c for c in course_analysis if c['enrolled']<x['students']*0.5]
    if whole_risk:
        risk_text="、".join(f"{c['name']}（挂科{c['failed']}人，挂科率{pct(c['failed']/c['enrolled'])}）" for c in whole_risk[:5])
        para(doc,f"整班课程中需重点关注{risk_text}。这些课程构成本班当前主要学业风险点，应结合课堂出勤、作业完成、阶段测验和课程难度进一步分析具体成因。")
    else: para(doc,"本班整班修读课程均已通过，未发现集中性课程挂科风险。建议继续保持过程性学习督导，并关注个别学生的学习波动。")
    if individual:
        individual_fail=sum(c['failed'] for c in individual)
        para(doc,f"另有{len(individual)}门重修、补修或个别修读课程，共{sum(c['enrolled'] for c in individual)}人次修读、{individual_fail}人次不及格。此类课程不宜直接作为整班教学质量判断依据，应按学生个人修读计划单独跟踪。")
    if top:
        doc.add_picture(str(charts[1]),width=Cm(15));caption(doc,"报告图2 课程挂科率TOP10（括号内为挂科人数）")
    else: para(doc,"本班无挂科课程，课程挂科率图略。")
    heading(doc,"五、整改措施")
    if whole_risk: para(doc,"1. 开展课程专项补强。围绕"+"、".join(c['name'] for c in whole_risk[:3])+"，联合任课教师梳理高频知识难点，组织专题答疑、阶段测试和学习任务清单，并以周为单位反馈完成情况。")
    else: para(doc,"1. 巩固课程学习成效。继续保持课堂考勤、作业反馈和阶段测验联动，及时识别成绩波动或学习投入下降的学生。")
    if x['fail_students']:
        para(doc,"2. 实施重点学生精准帮扶。对多门不及格学生建立“一人一档”，逐门明确补考、重修和学分修读计划，至少完成一次深度谈话，并根据需要启动朋辈帮扶、家校沟通或心理支持。")
    elif x['fail_students']: para(doc,"2. 对单门不及格学生逐人提醒补考安排，核实复习计划和课程困难，防止风险累积；补考后及时更新结果并闭环记录。")
    else: para(doc,"2. 对成绩临界学生开展预防性指导，帮助其形成稳定的复习计划和课程学习节奏。")
    para(doc,f"3. 做好分层提升。对60—69分段的{seg[3]}名学生开展学业预警和基础巩固，对70—79分段的{seg[2]}名学生强化方法指导，对80分以上的{good}名学生鼓励参与朋辈答疑和学习经验分享。")
    para(doc,f"4. 强化补考重修闭环。当前补考通过{makeup_pass}人次、待重修{pending}人次，应建立“通知—复习—考试—结果—学分”跟踪台账，避免遗漏报名、考试或重修选课节点。")
    para(doc,"5. 坚持动态监测。每月结合考勤、作业、课堂表现和谈心谈话记录更新重点关注名单；涉及明显心理压力、家庭困难或持续缺勤的学生，应按工作流程转介并协同支持。")
    doc.save(folder/f"{x['class']}.docx")

def process(upload, original, teacher_path=None, teacher_name=""):
    run_id=datetime.now().strftime("%Y%m%d-%H%M%S-")+uuid.uuid4().hex[:6]; out=RUNS/run_id;out.mkdir(); infile=out/"input.xlsx";infile.write_bytes(upload)
    headers, rows, records, errors=load_records(infile); classes=stats_for(records,rows)
    if not classes: raise ValueError("过滤后没有可统计的班级记录")
    teacher_audit={"provided": False, "rows": 0, "matched_classes": 0, "unmatched_classes": [], "outside_classes": []}
    if teacher_path:
        teachers, teacher_audit, teacher_issues = load_teachers(Path(teacher_path), [item["class"] for item in classes])
        errors.extend(teacher_issues)
        for item in classes:
            item["teacher"] = teachers.get(class_key(item["class"]), "")
        teacher_audit["source_file"] = teacher_name or Path(teacher_path).name
    else:
        for item in classes:
            item["teacher"] = ""
    sample={"year":classes[0]["courses"][0]["items"][0]["学年"],"semester":semester_name(classes[0]["courses"][0]["items"][0]["学期"]),"college":classes[0]["college"]}
    summary=out/"学业考试成绩分析统计表.xlsx";create_xlsx(classes,sample,summary); reports=out/"学业分析";reports.mkdir()
    for c in classes:create_report(c,sample,reports)
    archive=out/"学业分析报告.zip"
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED) as z:
        for f in reports.rglob("*.docx"):z.write(f,f.relative_to(reports))
    audit={"source_file":original,"teacher_list":teacher_audit,"rows":len(rows),"valid_records":len(records),"classes":len(classes),"students":len({r['学号'] for r in records}),"excluded_records":len(rows)-len(records)-len([item for item in errors if "row" in item and not item.get("source")]),"issues":errors,"created_at":datetime.now().isoformat()}
    (out/"校验报告.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding="utf-8")
    class_metrics=[]
    total_fail_records=0
    total_fail_students=0
    for index,item in enumerate(classes,1):
        fail_records=len(item["failures"]); fail_students=len(item["fail_students"])
        total_fail_records+=fail_records; total_fail_students+=fail_students
        class_metrics.append({
            "class_index":index,
            "students":item["students"],
            "course_count":len(item["courses"]),
            "exam_records":item["records"],
            "fail_records":fail_records,
            "fail_student_count":fail_students,
            "fail_rate":round(fail_records/item["records"],6) if item["records"] else 0,
            "fail_buckets":item["bucket"]
        })
    ai_summary={
        "privacy":"仅含匿名聚合统计，不含姓名、学号、班级名、专业名、课程名、身份证号及原始成绩",
        "classes":len(classes),"students":audit["students"],"valid_records":len(records),
        "excluded_records":audit["excluded_records"],"issue_count":len(errors),
        "fail_records":total_fail_records,"fail_student_class_counts":total_fail_students,
        "overall_fail_rate":round(total_fail_records/len(records),6) if records else 0,
        "class_metrics":class_metrics
    }
    return {"runId":run_id,"overview":audit,"aiSummary":ai_summary,"downloads":{"summary":f"/runs/{run_id}/{summary.name}","reports":f"/runs/{run_id}/{archive.name}","audit":f"/runs/{run_id}/校验报告.json"}}

class Handler(SimpleHTTPRequestHandler):
    def translate_path(self,path):
        p=unquote(urlparse(path).path)
        if p.startswith("/runs/"):return str(ROOT/p.lstrip("/"))
        if p=="/" or p=="/index.html":return str(STATIC/"index.html")
        return str(STATIC/p.lstrip("/"))
    def do_POST(self):
        if self.path!="/api/analyze":self.send_error(404);return
        try:
            content_type=self.headers.get("Content-Type","")
            match=re.search(r"boundary=([^;]+)",content_type)
            if not match: raise ValueError("无效的上传请求")
            boundary=match.group(1).strip('"').encode(); raw=self.rfile.read(int(self.headers.get("Content-Length","0")))
            part=next((x for x in raw.split(b"--"+boundary) if b'name="file"' in x),None)
            if not part or b"\r\n\r\n" not in part: raise ValueError("请选择Excel文件")
            header,data=part.split(b"\r\n\r\n",1); data=data.rsplit(b"\r\n",1)[0]
            found=re.search(br'filename="([^\"]*)"',header); name=Path((found.group(1).decode("utf-8","replace") if found else "input.xlsx")).name
            if Path(name).suffix.lower()!=".xlsx" or name.startswith("~$"):raise ValueError("仅支持非锁定的 .xlsx 文件")
            result=process(data,name);body=json.dumps({"ok":True,**result},ensure_ascii=False).encode()
            self.send_response(200);self.send_header("Content-Type","application/json; charset=utf-8");self.send_header("Content-Length",str(len(body)));self.end_headers();self.wfile.write(body)
        except Exception as e:
            body=json.dumps({"ok":False,"error":str(e)},ensure_ascii=False).encode();self.send_response(400);self.send_header("Content-Type","application/json; charset=utf-8");self.end_headers();self.wfile.write(body)

if __name__=="__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "--analyze":
        input_path = Path(sys.argv[2]).resolve()
        original_name = sys.argv[3] if len(sys.argv) >= 4 else input_path.name
        teacher_path = None
        teacher_name = ""
        if len(sys.argv) >= 6 and sys.argv[4] == "--teachers":
            teacher_path = Path(sys.argv[5]).resolve()
            teacher_name = sys.argv[6] if len(sys.argv) >= 7 else teacher_path.name
        try:
            result = process(input_path.read_bytes(), original_name, teacher_path, teacher_name)
            print(json.dumps({"ok": True, **result}, ensure_ascii=False))
        except Exception as error:
            print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False))
            raise SystemExit(2)
    else:
        port=int(sys.argv[1]) if len(sys.argv)>1 else 18818
        print(f"学业分析助手：http://127.0.0.1:{port}")
        ThreadingHTTPServer(("127.0.0.1",port),Handler).serve_forever()
